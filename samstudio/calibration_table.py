"""Extension-insensitive Excel calibration lookup, independent of COCO editing."""
from __future__ import annotations
import copy, math, unicodedata
from pathlib import Path
import numpy as np
from .domain import Calibration, LENGTH_UNITS


def image_stem(value):
    name = str(value).strip().replace('\\', '/').rsplit('/', 1)[-1]
    return unicodedata.normalize('NFC', Path(name).stem).casefold()


class PixelSizeTable:
    def __init__(self, entries, source, unit='nm'):
        self.entries, self.source, self.unit = entries, str(source), unit

    @classmethod
    def read(cls, path, sheet=None, unit='nm'):
        from openpyxl import load_workbook
        if unit not in LENGTH_UNITS:
            raise ValueError('Choose a supported physical unit per pixel.')
        book = load_workbook(path, read_only=True, data_only=True)
        try:
            rows = iter((book[sheet] if sheet else book.active).iter_rows(values_only=True))
            headers = next(rows, None)
            if not headers:
                raise ValueError('The workbook is empty.')
            columns = {str(v).strip().casefold(): i for i, v in enumerate(headers) if v is not None}
            if not {'converted_image_name', 'converted_size'} <= columns.keys():
                raise ValueError('Required columns: converted_image_name and converted_size.')
            entries = {}
            for number, row in enumerate(rows, 2):
                name, value = row[columns['converted_image_name']], row[columns['converted_size']]
                if name is None or not str(name).strip():
                    continue
                key = image_stem(name)
                if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value) or value <= 0:
                    raise ValueError(f'Excel row {number}: converted_size must be a positive number. Formulas need cached numeric results.')
                if key in entries and entries[key] != float(value):
                    raise ValueError(f'Excel row {number}: conflicting converted_size values for image stem {key}.')
                entries[key] = float(value)
            if not entries:
                raise ValueError('The workbook contains no valid image calibration rows.')
            return cls(entries, Path(path).resolve(), unit)
        finally:
            book.close()

    def calibration(self, name):
        value = self.entries.get(image_stem(name))
        return Calibration(value, value, self.unit, f'Excel converted_size: {self.source}', True) if value is not None else None


def apply_calibration(document, calibration):
    """Rebuild physical coordinates through any existing rotations."""
    source = copy.deepcopy(calibration)
    current = copy.deepcopy(calibration)
    for rotation in document.rotation:
        matrix = np.asarray(rotation['matrix'], float)
        basis = current.matrix() @ np.linalg.inv(matrix[:, :2])
        current.origin = (np.asarray(current.origin or [0., 0.]) - basis @ matrix[:, 2]).tolist()
        current.basis = basis.tolist()
        current.pixel_size_x, current.pixel_size_y = np.linalg.norm(basis, axis=0).tolist()
    if document.rotation:
        current.source += '; rotated coordinate basis'
    document.source_calibration, document.calibration = source, current


class CalibrationTableMixin:
    def _init_calibration_tables(self):
        self.pixel_size_table = None
        menu = self.statistics_menu
        menu.addSeparator()
        self.excel_calibration_action = self._action('Read pixel sizes from Excel…', self.pixel_table_dialog)
        menu.addAction(self.excel_calibration_action)
        menu.addAction(self._action('Forget automatic Excel lookup', self.forget_pixel_table))
        path = self.settings.value('pixel_size_table', '')
        if path:
            try:
                self.pixel_size_table = PixelSizeTable.read(path, self.settings.value('pixel_size_sheet', None), self.settings.value('pixel_size_unit', 'nm'))
            except Exception as error:
                self.statusBar().showMessage('Excel lookup was not loaded: ' + str(error), 15000)

    def forget_pixel_table(self):
        self.pixel_size_table = None
        for key in ('pixel_size_table', 'pixel_size_sheet', 'pixel_size_unit'):
            self.settings.remove(key)
        self.settings.sync()
        self.statusBar().showMessage('Automatic Excel lookup cleared. Existing image calibration is retained.', 10000)

    def pixel_table_dialog(self):
        from PySide6.QtWidgets import QFileDialog, QInputDialog
        from openpyxl import load_workbook
        path, _ = QFileDialog.getOpenFileName(self, 'Read converted_size from Excel', '', 'Excel workbooks (*.xlsx *.xlsm)')
        if not path or self.busy:
            return
        try:
            book = load_workbook(path, read_only=True); sheets = book.sheetnames; book.close()
            sheet, ok = QInputDialog.getItem(self, 'Pixel-size table', 'Worksheet', sheets, editable=False)
            if not ok:
                return
            unit, ok = QInputDialog.getItem(self, 'Pixel-size unit', 'Unit per pixel for converted_size', ['nm', 'µm', 'pm', 'Å', 'mm', 'm'], editable=False)
            if not ok:
                return
            lookup = PixelSizeTable.read(path, sheet, unit)
            matched, unmatched = 0, []
            for key, doc in self.documents.items():
                cal = lookup.calibration(doc.frame.source_path)
                if cal:
                    self.histories[key].apply('Excel pixel size', lambda d=doc, c=cal: apply_calibration(d, c))
                    matched += 1
                else:
                    unmatched.append(Path(doc.frame.source_path).name)
            self.pixel_size_table = lookup
            for key, value in [('pixel_size_table', str(Path(path).resolve())), ('pixel_size_sheet', sheet), ('pixel_size_unit', unit)]:
                self.settings.setValue(key, value)
            self.settings.sync(); self._refresh()
            self.statusBar().showMessage(f'Excel calibration applied to {matched} frames; {len(unmatched)} unmatched. The lookup also applies to later opens.', 15000)
        except Exception as error:
            self.notify('Pixel-size table', str(error))

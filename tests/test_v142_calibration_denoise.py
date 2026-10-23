"""Numerical and native GUI checks for the requested v1.4.2 workflows."""
import copy, json, threading
from pathlib import Path
import numpy as np
import pytest
from openpyxl import Workbook
from samstudio.calibration_table import PixelSizeTable, image_stem, apply_calibration
from samstudio.domain import Calibration, Document
from samstudio.image_processing import parameters, process_pixels
from samstudio.denoise_config import METHODS, parameter_html


def workbook(path, rows, headers=('converted_image_name','converted_size')):
    book=Workbook();book.active.append(headers)
    for row in rows:book.active.append(row)
    book.save(path);return path


@pytest.mark.parametrize('name',['sample.png','SAMPLE.TIFF',r'C:\data\Sample.dm3','/x/sample.emd','sample'])
def test_excel_extension_insensitive_lookup(tmp_path,name):
    lookup=PixelSizeTable.read(workbook(tmp_path/'pixels.xlsx',[('sample.png',.125)]))
    cal=lookup.calibration(name)
    assert cal.valid and cal.pixel_size_x==cal.pixel_size_y==.125 and cal.length_unit=='nm'
    assert lookup.calibration('other.tif') is None


@pytest.mark.parametrize('value',[0,-1,True,'0.2',None,'=1/5'])
def test_excel_invalid_calibration_rejected(tmp_path,value):
    with pytest.raises(ValueError,match='positive number'):
        PixelSizeTable.read(workbook(tmp_path/'pixels.xlsx',[('a.png',value)]))


def test_excel_conflicts_empty_headers_and_duplicates(tmp_path):
    with pytest.raises(ValueError,match='conflicting'):
        PixelSizeTable.read(workbook(tmp_path/'conflict.xlsx',[('a.png',1),('a.tif',2)]))
    with pytest.raises(ValueError,match='Required columns'):
        PixelSizeTable.read(workbook(tmp_path/'bad.xlsx',[('a.png',1)],('name','size')))
    table=PixelSizeTable.read(workbook(tmp_path/'valid.xlsx',[('a.png',1),('a.tiff',1)]))
    assert table.entries=={'a':1.}
    assert image_stem('particle.v2.tiff')=='particle.v2'


def test_excel_rotation_basis_and_native_recovery(document):
    from samstudio.transforms import rotate_document
    doc=document;rotate_document(doc,30,'rectangle');rotate_document(doc,45,'square')
    masks=[o.mask.copy() for o in doc.objects]
    apply_calibration(doc,Calibration(.4,.4,'nm','Excel',True))
    composed=np.eye(3)
    for rotation in doc.rotation:
        affine=np.eye(3);affine[:2]=rotation['matrix'];composed=affine@composed
    expected=.4*np.linalg.inv(composed)[:2,:]
    assert np.allclose(doc.calibration.matrix(),expected[:,:2])
    assert np.allclose(doc.calibration.origin,expected[:,2])
    assert all(np.array_equal(o.mask,m) for o,m in zip(doc.objects,masks))
    frame=copy.copy(doc.frame);frame.raw=doc.source_raw
    restored=Document(frame);restored.restore(doc.state())
    assert restored.source_calibration.pixel_size_x==.4 and restored.calibration==doc.calibration


def test_excel_gui_import_applies_existing_and_future_images(window,document,tmp_path,monkeypatch):
    from tests.test_v141_processing_workspace import mount
    from PySide6.QtWidgets import QFileDialog,QInputDialog
    from PIL import Image
    from samstudio.imageio import LoadResult,load_file
    mount(window,document)
    table=workbook(tmp_path/'pixels.xlsx',[('SOURCE.png',.15),('later.tiff',.25)])
    monkeypatch.setattr(QFileDialog,'getOpenFileName',lambda *a,**k:(str(table),''))
    answers=iter([('Sheet',True),('nm',True)])
    monkeypatch.setattr(QInputDialog,'getItem',lambda *a,**k:next(answers))
    window.pixel_table_dialog();assert not window.errors
    assert document.calibration.pixel_size_x==.15
    restored=Document(document.frame);window.stores[document.frame.key].load(restored)
    assert restored.calibration.pixel_size_x==.15
    Image.new('L',(32,32),127).save(tmp_path/'later.png')
    window._opened(LoadResult(load_file(tmp_path/'later.png')))
    assert window.document.calibration.pixel_size_x==.25
    window.forget_pixel_table();assert window.pixel_size_table is None
    assert window.document.calibration.pixel_size_x==.25


def test_dot_disables_single_and_roi_including_programmatic_selection(window,document):
    from tests.test_v141_processing_workspace import mount
    mount(window,document);window.mode.setCurrentText('Segment')
    for scope in ('single','roi'):
        window.tool_actions['rectangle'].trigger();window.scope.setCurrentIndex(window.scope.findData(scope))
        window.tool_actions['point'].trigger();assert window.scope.currentData()=='prompt'
        index=window.scope.findData(scope)
        assert not window.scope.model().item(index).isEnabled()
        window.scope.setCurrentIndex(index);assert window.scope.currentData()=='prompt'
    window.tool_actions['ellipse'].trigger()
    assert all(window.scope.model().item(window.scope.findData(v)).isEnabled() for v in ('single','roi'))


def notebook_reference(x, method, p):
    """Independent direct translation of the supplied notebook, one RGB plane."""
    from scipy.ndimage import distance_transform_edt
    from scipy.fft import dctn,idctn
    flat=np.zeros(x.size,bool)
    flat[np.random.default_rng(p['random_seed']).choice(x.size,max(1,round(p['sample_fraction']*x.size)),replace=False)]=True
    mask=flat.reshape(x.shape);idx=distance_transform_edt(~mask,return_distances=False,return_indices=True)
    estimate=x[tuple(idx)].copy()
    for _ in range(p['iterations']):
        if method=='DCT sparsity':
            coeff=dctn(estimate,norm='ortho',axes=(-2,-1))
            keep=max(1,round(p['sparsity_fraction']*x.size))
            threshold=np.partition(abs(coeff).ravel(),-keep)[-keep]
            proposed=idctn(coeff*(abs(coeff)>=threshold),norm='ortho',axes=(-2,-1))
        else:
            gx=np.diff(estimate,axis=1,append=estimate[:,-1:]);gy=np.diff(estimate,axis=0,append=estimate[-1:,:])
            mag=np.sqrt(gx*gx+gy*gy+p['gradient_epsilon']);dx=gx/mag;dy=gy/mag
            div=dx.copy();div[:,1:]-=dx[:,:-1];div+=dy;div[1:]-=dy[:-1]
            proposed=np.clip(estimate+p['tv_step']*div,0,1)
        residual=x[mask]-proposed[mask]
        update=residual if p['data_loss']=='l2' else p['huber_delta']*np.sign(residual) if p['data_loss']=='l1' else np.clip(residual,-p['huber_delta'],p['huber_delta'])
        proposed[mask]+=update;estimate=np.clip(proposed,0,1)
    return estimate


@pytest.mark.parametrize('method',['DCT sparsity','TV iterative (notebook)'])
@pytest.mark.parametrize('loss',['l2','l1','huber'])
def test_sampled_reconstruction_matches_notebook_exactly(method,loss):
    from samstudio.compressed_sensing import iterative_restore
    x=np.random.default_rng(3).random((23,29),dtype=np.float32)
    p=parameters(method,{'task':'reconstruct','data_loss':loss,**({'tv_step':.08} if method.startswith('TV') else {})})
    assert np.array_equal(iterative_restore(x,method,p),notebook_reference(x,method,p))


@pytest.mark.parametrize('method',['DCT sparsity','TV iterative (notebook)'])
def test_hard_full_sampling_is_identity_but_denoise_changes_input(method):
    x=np.random.default_rng(2).random((32,40),dtype=np.float32)
    restored=process_pixels(x,method,{'task':'reconstruct','sample_fraction':1})
    assert np.allclose(restored,x,atol=1e-7)
    assert np.mean(abs(process_pixels(x,method)-x))>.001


@pytest.mark.parametrize('method',['DCT sparsity','TV iterative (notebook)','BM3D','Sparse dictionary learning'])
def test_actual_backends_reduce_known_gaussian_phantom_error(method):
    yy,xx=np.mgrid[:64,:64];clean=np.full((64,64),.25);clean[(xx-32)**2+(yy-32)**2<144]=.7
    noisy=np.clip(clean+np.random.default_rng(7).normal(0,.06,clean.shape),0,1).astype(np.float32)
    p={'sigma':.06} if method=='BM3D' else {'dictionary_iterations':4,'coding_iterations':15,'training_patches':256,'dictionary_atoms':16} if method=='Sparse dictionary learning' else {}
    result=process_pixels(noisy,method,p)
    assert np.mean((result-clean)**2)<np.mean((noisy-clean)**2)


@pytest.mark.parametrize('method',list(METHODS))
def test_every_control_has_default_explanation_and_recommendation(method):
    for key,spec in METHODS[method].items():
        assert spec['help'] and spec['recommended'] and key in parameters(method)
        page=parameter_html(method,key)
        assert 'Default:' in page and 'Suggested start:' in page


def test_processing_dialog_notebook_preset_and_defaults(qtbot):
    from samstudio.processing_ui import ProcessingDialog
    dialog=ProcessingDialog(['DCT sparsity','TV iterative (notebook)','Noise2Void']);qtbot.addWidget(dialog)
    assert dialog.values()['task']=='denoise' and not dialog.controls['sample_fraction'].isEnabled()
    dialog.notebook_preset();p=dialog.values()
    assert p['task']=='reconstruct' and p['max_side']==256 and p['random_seed']==7 and p['iterations']==40
    assert not dialog.controls['fidelity_step'].isEnabled()
    dialog.method.setCurrentText('TV iterative (notebook)');assert dialog.values()['tv_step']==.01
    dialog.notebook_preset();assert dialog.values()['tv_step']==.08
    dialog.reset_controls();assert dialog.values()['task']=='denoise'
    dialog.method.setCurrentText('Noise2Void');assert not dialog.controls['model_path'].isEnabled()
    dialog.controls['run_mode'].setCurrentText('predict');assert dialog.controls['model_path'].isEnabled()
    assert not dialog.controls['epochs'].isEnabled()


@pytest.mark.parametrize('method',['Noise2Void','N2V2'])
def test_real_blind_spot_train_save_reload_and_variant(tmp_path,method):
    from samstudio.inference import initialize_torch_runtime
    from samstudio.neural_denoising import neural_denoise
    initialize_torch_runtime()
    x=np.random.default_rng(1).random((64,64),dtype=np.float32)
    p=parameters(method,dict(epochs=1,steps_per_epoch=1,patch_size=32,batch_size=2,depth=1,initial_filters=8,
        validation_patches=2,tile_size=32,tile_overlap=4,output_directory=str(tmp_path)))
    y=neural_denoise(x,method,p)
    assert y.shape==x.shape and np.isfinite(y).all() and not np.array_equal(y,x)
    record=json.loads((tmp_path/'channel_0/samstudio_model.json').read_text())
    assert record['method']==method and (tmp_path/'channel_0/samstudio_model.ckpt').stat().st_size>1024
    p.update(run_mode='predict',model_path=str(tmp_path))
    z=neural_denoise(x,method,p);assert np.array_equal(y,z)
    with pytest.raises(ValueError,match='variant'):
        neural_denoise(x,'N2V2' if method=='Noise2Void' else 'Noise2Void',p)


@pytest.mark.parametrize('method',['DCT sparsity','TV iterative (notebook)','BM3D','Sparse dictionary learning','Noise2Void','N2V2'])
def test_cancelled_backend_does_not_return_image(method):
    event=threading.Event();event.set()
    with pytest.raises(InterruptedError):process_pixels(np.ones((64,64),np.float32)*.3,method,cancel=event)


@pytest.mark.parametrize('shape',[(1,1),(8,8),(5,19)])
def test_bm3d_small_images_use_safe_native_context(shape):
    raw=np.full(shape,.3,np.float32)
    out=process_pixels(raw,'BM3D')
    assert out.shape==shape and np.isfinite(out).all() and np.allclose(out,raw,atol=1e-4)


@pytest.mark.parametrize('method',['DCT sparsity','Noise2Void'])
def test_new_processing_worker_preview_apply_undo_and_model_persistence(window,document,tmp_path,monkeypatch,qtbot,method):
    from tests.test_v141_processing_workspace import mount
    from PySide6.QtWidgets import QDialog
    from PySide6.QtCore import QTimer
    from samstudio.processing_ui import ProcessingDialog
    from samstudio.dialogs import ImagePreviewDialog
    mount(window,document);original=document.state();previews=[];ticks=[]
    p=parameters(method,dict(epochs=1,steps_per_epoch=1,patch_size=32,batch_size=2,depth=1,initial_filters=8,
        validation_patches=2,tile_size=32,tile_overlap=4,output_directory=str(tmp_path/'trained'))) if method=='Noise2Void' else parameters(method)
    monkeypatch.setattr(ProcessingDialog,'exec',lambda self:QDialog.DialogCode.Accepted)
    monkeypatch.setattr(ProcessingDialog,'values',lambda self:p)
    monkeypatch.setattr(ImagePreviewDialog,'exec',lambda self:(previews.append(self.windowTitle()) or QDialog.DialogCode.Accepted))
    timer=QTimer();timer.timeout.connect(lambda:ticks.append(1));timer.start(5)
    window.process_image_dialog([method]);qtbot.waitUntil(lambda:not window.busy,timeout=60000);timer.stop()
    assert previews==['Preview '+method] and not window.errors and document.working is not None
    assert [o.to_dict() for o in document.foreground_objects()]==original['objects']
    assert ticks
    window.undo();assert document.working is None and document.state()['objects']==original['objects']
    if method=='Noise2Void':assert (tmp_path/'trained/channel_0/samstudio_model.ckpt').is_file()
